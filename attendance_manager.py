# /DigitalFaceAttendanceSystem/attendance_manager.py

import openpyxl
import pickle
import cv2
import face_recognition
import numpy as np
import os
import re 
from pathlib import Path
from datetime import datetime, timedelta
import csv
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font
import shutil 
import sqlite3
import pandas as pd

# --- Configuration ---
BASE_DIR = Path(__file__).resolve().parent
DATASET_DIR = BASE_DIR / "dataset"
STUDENTS_DIR = BASE_DIR / "students"
ENCODINGS_DIR = BASE_DIR / "encodings"
UNKNOWN_LOGS_DIR = BASE_DIR / "unknown_logs"
ENCODINGS_PATH = ENCODINGS_DIR / "encodings.pkl"
DB_FILE = BASE_DIR / "attendance.db" 

# --- Helper Function: Database connection ---
def create_connection():
    try:
        conn = sqlite3.connect(DB_FILE)
        conn.execute("PRAGMA foreign_keys = 1")
        return conn
    except Exception as e:
        print(f"Database connection error: {e}")
        return None

# --- Environment Setup ---
def setup_environment():
    DATASET_DIR.mkdir(parents=True, exist_ok=True)
    STUDENTS_DIR.mkdir(parents=True, exist_ok=True)
    ENCODINGS_DIR.mkdir(parents=True, exist_ok=True)
    UNKNOWN_LOGS_DIR.mkdir(parents=True, exist_ok=True)

    if not ENCODINGS_PATH.exists():
        with open(ENCODINGS_PATH, 'wb') as f:
            pickle.dump({"encodings": [], "data": []}, f)
            
    conn = None
    try:
        conn = create_connection()
        cursor = conn.cursor()
        cursor.execute("""
        CREATE TABLE IF NOT EXISTS students (
            roll_no TEXT PRIMARY KEY, name TEXT NOT NULL, department TEXT,
            class_name TEXT, prn_no TEXT, contact_no TEXT, dob TEXT
        );""")
        cursor.execute("""
        CREATE TABLE IF NOT EXISTS attendance_logs (
            log_id INTEGER PRIMARY KEY AUTOINCREMENT, roll_no TEXT NOT NULL,
            name TEXT, date TEXT, timestamp TEXT, lecture_session TEXT,
            FOREIGN KEY (roll_no) REFERENCES students (roll_no) ON DELETE CASCADE
        );""")
        cursor.execute("""
        CREATE TABLE IF NOT EXISTS activity_log (
            log_id INTEGER PRIMARY KEY AUTOINCREMENT, timestamp TEXT,
            username TEXT, action TEXT
        );""")
        cursor.execute("""
        CREATE TABLE IF NOT EXISTS users (
            username TEXT PRIMARY KEY, password_hash TEXT NOT NULL,
            role TEXT NOT NULL, full_name TEXT, contact_no TEXT,
            subject_specialization TEXT, failed_attempts INTEGER DEFAULT 0,
            lockout_until TEXT, department TEXT
        );""")
        cursor.execute("CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT);")
        conn.commit()
    except Exception as e:
        if "duplicate column name" not in str(e) and "table users has no column named department" not in str(e):
            try:
                cursor.execute("ALTER TABLE users ADD COLUMN failed_attempts INTEGER DEFAULT 0")
                cursor.execute("ALTER TABLE users ADD COLUMN lockout_until TEXT")
                cursor.execute("ALTER TABLE users ADD COLUMN department TEXT")
                conn.commit()
                print("Database 'users' table updated with new columns.")
            except Exception as e_alter:
                if "duplicate column name" not in str(e_alter):
                    print(f"Error updating database schema: {e_alter}")
    finally:
        if conn:
            conn.close()

# --- Activity Log Function ---
def log_activity(username, action):
    conn = None
    try:
        conn = create_connection()
        cursor = conn.cursor()
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        cursor.execute("INSERT INTO activity_log (timestamp, username, action) VALUES (?, ?, ?)", (timestamp, username, action))
        conn.commit()
    except Exception as e:
        print(f"Error logging activity: {e}")
    finally:
        if conn:
            conn.close()

# --- Encoding Management ---
def load_encodings():
    try:
        with open(ENCODINGS_PATH, 'rb') as f:
            data = pickle.load(f)
        return data.get("encodings", []), data.get("data", [])
    except (FileNotFoundError, EOFError, pickle.UnpicklingError):
        return [], []
def save_encodings(encodings, data):
    with open(ENCODINGS_PATH, 'wb') as f:
        pickle.dump({"encodings": encodings, "data": data}, f)

# --- Student Management ---
def add_student(roll_no, name, dob, class_name, department, prn_no, contact_no, images):
    if not all([roll_no, name, dob, class_name, department, prn_no, contact_no, images]):
        raise ValueError("All fields and at least one image are required.")
    student_dir = DATASET_DIR / f"{roll_no}_{name}"
    student_dir.mkdir(parents=True, exist_ok=True)
    for i, img in enumerate(images):
        img_path = student_dir / f"{i+1}.png"
        cv2.imwrite(str(img_path), img)
    conn = None
    try:
        conn = create_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM students WHERE roll_no = ?", (roll_no,))
        if cursor.fetchone():
            raise ValueError(f"Student with Roll No {roll_no} already exists.")
        cursor.execute("""
        INSERT INTO students (roll_no, name, department, class_name, prn_no, contact_no, dob)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (roll_no, name, department, class_name, prn_no, contact_no, dob))
        conn.commit()
    except sqlite3.IntegrityError:
         raise ValueError(f"Student with Roll No {roll_no} already exists.")
    except Exception as e:
        raise Exception(f"Failed to update database: {e}")
    finally:
        if conn:
            conn.close()
    try:
        known_encodings, known_data = load_encodings()
        for img_from_cam in images:
            if img_from_cam is None: continue
            if img_from_cam.dtype != np.uint8:
                img_from_cam = img_from_cam.astype(np.uint8)
            img_rgb = None
            if len(img_from_cam.shape) == 2:
                img_rgb = cv2.cvtColor(img_from_cam, cv2.COLOR_GRAY2RGB)
            elif len(img_from_cam.shape) == 3 and img_from_cam.shape[2] == 3:
                img_rgb = cv2.cvtColor(img_from_cam, cv2.COLOR_BGR2RGB)
            elif len(img_from_cam.shape) == 3 and img_from_cam.shape[2] == 4:
                img_rgb = cv2.cvtColor(img_from_cam, cv2.COLOR_BGRA2RGB)
            else:
                raise Exception(f"Unsupported image shape: {img_from_cam.shape}")
            boxes = face_recognition.face_locations(img_rgb, model='hog')
            new_encodings = face_recognition.face_encodings(img_rgb, boxes)
            if new_encodings:
                known_encodings.append(new_encodings[0])
                known_data.append({"roll": roll_no, "name": name})
            else:
                raise ValueError(f"No face detected in the captured image for {name}.")
        save_encodings(known_encodings, known_data)
    except Exception as e:
        try:
            conn = create_connection()
            conn.execute("DELETE FROM students WHERE roll_no = ?", (roll_no,))
            conn.commit()
        except Exception as rb_e:
            print(f"ROLLBACK FAILED: {rb_e}")
        finally:
            if conn:
                conn.close()
        raise Exception(f"Failed to generate or save encodings: {e}")
def mark_attendance(roll_no, name, session_id):
    conn = None
    try:
        conn = create_connection()
        cursor = conn.cursor()
        today = datetime.now().strftime("%Y-%m-%d")
        timestamp = datetime.now().strftime("%H:%M:%S")
        cursor.execute("SELECT * FROM attendance_logs WHERE roll_no = ? AND lecture_session = ?", (roll_no, session_id))
        if cursor.fetchone():
            return False 
        cursor.execute("""
        INSERT INTO attendance_logs (roll_no, name, date, timestamp, lecture_session)
        VALUES (?, ?, ?, ?, ?)
        """, (roll_no, name, today, timestamp, session_id))
        conn.commit()
        return True
    except Exception as e:
        print(f"Error marking attendance: {e}")
        return False
    finally:
        if conn:
            conn.close()
def get_student_details_dict(department=None):
    student_data = {}
    conn = None
    try:
        conn = create_connection()
        cursor = conn.cursor()
        query = "SELECT * FROM students"
        params = []
        if department:
            query += " WHERE department = ?"
            params.append(department)
        cursor.execute(query, params)
        for row in cursor.fetchall():
            roll_no_str = str(row[0])
            student_data[roll_no_str] = {
                "roll_no": row[0], "name": row[1], "department": row[2], 
                "class_name": row[3], "prn_no": row[4], "contact_no": row[5], 
                "dob": row[6]
            }
        return student_data
    except Exception as e:
        print(f"Error loading student details: {e}")
        return {}
    finally:
        if conn:
            conn.close()
def get_present_students_for_subject(subject_name, department=None):
    present_students = {}
    conn = None
    try:
        conn = create_connection()
        cursor = conn.cursor()
        query = """
        SELECT T1.roll_no, T1.date, T1.timestamp 
        FROM attendance_logs AS T1
        INNER JOIN students AS T2 ON T1.roll_no = T2.roll_no
        WHERE T1.lecture_session = ?
        """
        params = [subject_name]
        if department:
            query += " AND T2.department = ?"
            params.append(department)
        cursor.execute(query, params)
        for row in cursor.fetchall():
            present_students[str(row[0])] = (row[1], row[2])
        return present_students
    except Exception as e:
        print(f"Error loading attendance log: {e}")
        return {}
    finally:
        if conn:
            conn.close()
def generate_subject_summary_report(subject_name, department=None):
    all_students = get_student_details_dict(department)
    present_students = get_present_students_for_subject(subject_name, department)
    report_data, stats = [], {}
    total_present, total_absent, total_students = 0, 0, len(all_students)
    for roll_no, details in all_students.items():
        if roll_no in present_students:
            status, (date, timestamp) = "P", present_students[roll_no]
            total_present += 1
        else:
            status, date, timestamp = "A", "N/A", "N/A"
            total_absent += 1
        report_data.append([
            roll_no, details["name"], status, date, timestamp,
            details["department"], details["class_name"],
            details["prn_no"], details["contact_no"]
        ])
    present_percentage = (total_present / total_students) * 100 if total_students > 0 else 0
    stats = {
        "total": total_students, "present": total_present,
        "absent": total_absent, "percentage": f"{present_percentage:.2f}%"
    }
    return report_data, stats
def export_subject_report_to_excel(filepath, report_data, stats, subject_name):
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Report for {subject_name}"
        green_font, red_font = Font(color="008000", bold=True), Font(color="FF0000", bold=True)
        ws.append([f"Attendance Report for Subject: {subject_name}"])
        ws.append([])
        ws.append(["Total Students:", stats["total"]])
        ws.append(["Total Present:", stats["present"]])
        ws.append(["Total Absent:", stats["absent"]])
        ws.append(["Present Percentage:", stats["percentage"]])
        ws.append([])
        headers = ["Roll No", "Name", "Status", "Date", "Timestamp", "Department", "Class", "PRN No", "Contact No"]
        ws.append(headers)
        for row in report_data:
            ws.append(row)
            cell = ws.cell(row=ws.max_row, column=3)
            if row[2] == "P": cell.font = green_font
            elif row[2] == "A": cell.font = red_font
        for i, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = len(header) + 10
        wb.save(filepath)
        return True
    except Exception as e:
        print(f"Error exporting to Excel: {e}")
        return False
def get_sessions_for_date(date_str, department=None):
    sessions = set()
    conn = None
    try:
        conn = create_connection()
        cursor = conn.cursor()
        query = """
        SELECT DISTINCT T1.lecture_session 
        FROM attendance_logs AS T1
        INNER JOIN students AS T2 ON T1.roll_no = T2.roll_no
        WHERE T1.date = ?
        """
        params = [date_str]
        if department:
            query += " AND T2.department = ?"
            params.append(department)
        cursor.execute(query, params)
        for row in cursor.fetchall():
            sessions.add(row[0])
        return sorted(list(sessions))
    except Exception as e:
        print(f"Error loading sessions: {e}")
        return []
    finally:
        if conn:
            conn.close()
def get_attendance_by_date(date_str, department=None):
    attendance_data = {}
    conn = None
    try:
        conn = create_connection()
        cursor = conn.cursor()
        query = """
        SELECT T1.roll_no, T1.lecture_session 
        FROM attendance_logs AS T1
        INNER JOIN students AS T2 ON T1.roll_no = T2.roll_no
        WHERE T1.date = ?
        """
        params = [date_str]
        if department:
            query += " AND T2.department = ?"
            params.append(department)
        cursor.execute(query, params)
        for row in cursor.fetchall():
            roll_no, session_id = str(row[0]), str(row[1])
            attendance_data[(roll_no, session_id)] = True
        return attendance_data
    except Exception as e:
        print(f"Error loading attendance log: {e}")
        return {}
    finally:
        if conn:
            conn.close()
def generate_date_summary_report(date_str, department=None):
    all_students = get_student_details_dict(department)
    sessions = get_sessions_for_date(date_str, department)
    attendance_marks = get_attendance_by_date(date_str, department)
    headers = ["Roll No", "Name", "Department", "PRN No"]
    headers.extend(sessions)
    final_report_data = []
    session_stats = {session: {"present": 0, "absent": 0} for session in sessions}
    for roll_no, details in all_students.items():
        row_data = [
            roll_no, details["name"], details["department"], details["prn_no"]
        ]
        for session_id in sessions:
            if (roll_no, session_id) in attendance_marks:
                row_data.append("P")
                session_stats[session_id]["present"] += 1
            else:
                row_data.append("A")
                session_stats[session_id]["absent"] += 1
        final_report_data.append(row_data)
    final_stats = {}
    total_students = len(all_students)
    for session_name, counts in session_stats.items():
        p, a = counts["present"], counts["absent"]
        perc = f"{(p / total_students) * 100:.2f}%" if total_students > 0 else "0.00%"
        final_stats[session_name] = {"present": p, "absent": a, "percentage": perc, "total": total_students}
    return headers, final_report_data, final_stats
def export_date_summary_to_excel(filepath, headers, report_data, stats, date_str):
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Attendance Report {date_str}"
        green_font, red_font = Font(color="008000", bold=True), Font(color="FF0000", bold=True)
        ws.append([f"Attendance Report for Date: {date_str}"])
        ws.append([])
        ws.append(headers)
        for row in report_data:
            ws.append(row)
            for i, value in enumerate(row):
                if value == "P":
                    ws.cell(row=ws.max_row, column=i+1).font = green_font
                elif value == "A":
                    ws.cell(row=ws.max_row, column=i+1).font = red_font
        ws.append([])
        ws.append(["--- Session Summary ---"])
        ws.append(["Session", "Total", "Present", "Absent", "Percentage"])
        for session_name, stat in stats.items():
            ws.append([session_name, stat['total'], stat['present'], stat['absent'], stat['percentage']])
        for i, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = len(header) + 5
        wb.save(filepath)
        return True
    except Exception as e:
        print(f"Error exporting to Excel: {e}")
        return False
def get_all_students_list(department=None):
    students_list = []
    conn = None
    try:
        conn = create_connection()
        cursor = conn.cursor()
        query = "SELECT * FROM students"
        params = []
        if department:
            query += " WHERE department = ?"
            params.append(department)
        try:
            cursor.execute(f"{query} ORDER BY CAST(roll_no AS INTEGER)", params)
        except sqlite3.Error:
            cursor.execute(f"{query} ORDER BY roll_no", params)
        for row in cursor.fetchall():
            students_list.append(list(row))
        return students_list
    except Exception as e:
        print(f"Error reading student records: {e}")
        return []
    finally:
        if conn:
            conn.close()
def delete_student_by_roll(roll_no_to_delete):
    roll_no_to_delete = str(roll_no_to_delete)
    student_name = ""
    conn = None
    try:
        conn = create_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT name FROM students WHERE roll_no = ?", (roll_no_to_delete,))
        result = cursor.fetchone()
        if result:
            student_name = result[0]
        else:
            raise ValueError(f"Student with Roll No {roll_no_to_delete} not found in Database.")
        cursor.execute("DELETE FROM students WHERE roll_no = ?", (roll_no_to_delete,))
        conn.commit()
    except Exception as e:
        raise Exception(f"Error deleting from Database: {e}")
    finally:
        if conn:
            conn.close()
    try:
        if ENCODINGS_PATH.exists():
            known_encodings, known_data = load_encodings()
            new_encodings = [enc for enc, data in zip(known_encodings, known_data) if str(data.get("roll")) != roll_no_to_delete]
            new_data = [data for data in known_data if str(data.get("roll")) != roll_no_to_delete]
            save_encodings(new_encodings, new_data)
    except Exception as e:
        raise Exception(f"Error updating encodings.pkl: {e}")
    try:
        if student_name:
            folder_path_gui = DATASET_DIR / f"{roll_no_to_delete}_{student_name}"
            if folder_path_gui.exists():
                shutil.rmtree(folder_path_gui)
            sanitized_name = student_name.replace(' ', '_')
            sanitized_name = ''.join(e for e in sanitized_name if e.isalnum() or e == '_' or e == '-')
            folder_path_console = STUDENTS_DIR / sanitized_name
            if folder_path_console.exists():
                shutil.rmtree(folder_path_console)
    except Exception as e:
        print(f"Warning: Could not delete dataset folder for {roll_no_to_delete}: {e}")
    return True
def get_student_by_roll(roll_no):
    conn = None
    try:
        conn = create_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM students WHERE roll_no = ?", (str(roll_no),))
        row = cursor.fetchone()
        if row:
            student_data = {
                "roll_no": row[0], "name": row[1], "department": row[2],
                "class_name": row[3], "prn_no": row[4], "contact_no": row[5],
                "dob": row[6]
            }
            return student_data
        return None 
    except Exception as e:
        print(f"Error getting student by roll: {e}")
        return None
    finally:
        if conn:
            conn.close()
def update_student_data(roll_no, new_data):
    conn = None
    try:
        conn = create_connection()
        cursor = conn.cursor()
        cursor.execute("""
        UPDATE students 
        SET department = ?, class_name = ?, prn_no = ?, contact_no = ?, dob = ?
        WHERE roll_no = ?
        """, (
            new_data["department"], new_data["class_name"], new_data["prn_no"],
            new_data["contact_no"], new_data["dob"], str(roll_no)
        ))
        conn.commit()
        if cursor.rowcount == 0:
             raise ValueError(f"Student with Roll No {roll_no} not found.")
        return True
    except Exception as e:
        raise Exception(f"Error updating student data: {e}")
    finally:
        if conn:
            conn.close()
def generate_individual_report(roll_no, department=None):
    roll_no = str(roll_no)
    student_details = get_student_by_roll(roll_no)
    if not student_details:
        raise ValueError(f"Student with Roll No {roll_no} not found.")
    if department and student_details.get("department") != department:
        raise ValueError(f"Aap is student (Roll No: {roll_no}) ka report dekhne ke liye authorized nahi hain.")
    all_sessions_in_system = set()
    present_sessions = set()
    conn = None
    try:
        conn = create_connection()
        cursor = conn.cursor()
        query = "SELECT DISTINCT T1.date, T1.lecture_session FROM attendance_logs AS T1"
        params = []
        if department:
            query += " INNER JOIN students AS T2 ON T1.roll_no = T2.roll_no WHERE T2.department = ?"
            params.append(department)
        cursor.execute(query, params)
        for row in cursor.fetchall():
            all_sessions_in_system.add((str(row[0]), str(row[1])))
        cursor.execute("SELECT date, lecture_session FROM attendance_logs WHERE roll_no = ?", (roll_no,))
        for row in cursor.fetchall():
            present_sessions.add((str(row[0]), str(row[1])))
    except Exception as e:
        raise Exception(f"Error reading attendance records: {e}")
    finally:
        if conn:
            conn.close()
    report_data = []
    total_lectures = len(all_sessions_in_system)
    total_present = len(present_sessions)
    total_absent = total_lectures - total_present
    sorted_sessions = sorted(list(all_sessions_in_system), key=lambda x: x[0])
    for date, subject in sorted_sessions:
        status = "P" if (date, subject) in present_sessions else "A"
        report_data.append([subject, date, status])
    percentage = (total_present / total_lectures) * 100 if total_lectures > 0 else 0
    stats = {
        "total": total_lectures, "present": total_present,
        "absent": total_absent, "percentage": f"{percentage:.2f}%"
    }
    return student_details, report_data, stats
def get_dashboard_stats(department=None):
    stats = {"total": 0, "present": 0, "absent": 0}
    conn = None
    try:
        conn = create_connection()
        cursor = conn.cursor()
        query = "SELECT COUNT(*) FROM students"
        params = []
        if department:
            query += " WHERE department = ?"
            params.append(department)
        cursor.execute(query, params)
        stats["total"] = cursor.fetchone()[0]
        today = datetime.now().strftime("%Y-%m-%d")
        query = """
        SELECT COUNT(DISTINCT T1.roll_no) 
        FROM attendance_logs AS T1
        INNER JOIN students AS T2 ON T1.roll_no = T2.roll_no
        WHERE T1.date = ?
        """
        params = [today]
        if department:
            query += " AND T2.department = ?"
            params.append(department)
        cursor.execute(query, params)
        stats["present"] = cursor.fetchone()[0]
        stats["absent"] = stats["total"] - stats["present"]
        return stats
    except Exception as e:
        print(f"Error getting dashboard stats: {e}")
        return stats
    finally:
        if conn:
            conn.close()
def add_photos_to_student(roll_no, name, images):
    if not all([roll_no, name, images]):
        raise ValueError("Roll No, Name, aur images zaroori hain.")
    student_dir = DATASET_DIR / f"{roll_no}_{name}"
    student_dir.mkdir(parents=True, exist_ok=True) 
    try:
        for img in images:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
            img_path = student_dir / f"update_{timestamp}.png"
            cv2.imwrite(str(img_path), img)
    except Exception as e:
        raise Exception(f"Nayi images save karne mein error: {e}")
    new_encodings_list = []
    for img_from_cam in images:
        if img_from_cam is None: continue
        if img_from_cam.dtype != np.uint8:
            img_from_cam = img_from_cam.astype(np.uint8)
        img_rgb = None
        if len(img_from_cam.shape) == 2:
            img_rgb = cv2.cvtColor(img_from_cam, cv2.COLOR_GRAY2RGB)
        elif len(img_from_cam.shape) == 3 and img_from_cam.shape[2] == 3:
            img_rgb = cv2.cvtColor(img_from_cam, cv2.COLOR_BGR2RGB)
        elif len(img_from_cam.shape) == 3 and img_from_cam.shape[2] == 4:
            img_rgb = cv2.cvtColor(img_from_cam, cv2.COLOR_BGRA2RGB)
        else:
            continue 
        boxes = face_recognition.face_locations(img_rgb, model='hog')
        new_encodings = face_recognition.face_encodings(img_rgb, boxes)
        if new_encodings:
            new_encodings_list.append(new_encodings[0])
        else:
            raise ValueError(f"Aapki di gayi nayi photo(s) mein koi chehra detect nahi hua.")
    if not new_encodings_list:
        raise ValueError("Nayi photos mein se koi bhi face detect nahi hua.")
    try:
        known_encodings, known_data = load_encodings()
        for new_enc in new_encodings_list:
            known_encodings.append(new_enc)
            known_data.append({"roll": roll_no, "name": name})
        save_encodings(known_encodings, known_data)
        return True
    except Exception as e:
        raise Exception(f"Nayi encodings save karne mein error: {e}")
def export_individual_report_to_excel(filepath, student_details, report_data, stats):
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Report for {student_details['roll_no']}"
        green_font, red_font = Font(color="008000", bold=True), Font(color="FF0000", bold=True)
        ws.append(["Student Report Card"])
        ws.append(["Name:", student_details["name"]])
        ws.append(["Roll No:", student_details["roll_no"]])
        ws.append(["Department:", student_details["department"]])
        ws.append(["Class:", student_details["class_name"]])
        ws.append(["PRN No:", student_details["prn_no"]])
        ws.append([]) 
        ws.append(["Attendance Summary"])
        ws.append(["Total Lectures:", stats["total"]])
        ws.append(["Total Present:", stats["present"]])
        ws.append(["Total Absent:", stats["absent"]])
        ws.append(["Percentage:", stats["percentage"]])
        ws.append([]) 
        headers = ["Subject", "Date", "Status"]
        ws.append(headers)
        for row in report_data:
            ws.append(row)
            status = row[2] 
            cell = ws.cell(row=ws.max_row, column=3)
            if status == "P":
                cell.font = green_font
            elif status == "A":
                cell.font = red_font
        ws.column_dimensions[get_column_letter(1)].width = 25
        ws.column_dimensions[get_column_letter(2)].width = 15
        ws.column_dimensions[get_column_letter(3)].width = 10
        wb.save(filepath)
        return True
    except Exception as e:
        print(f"Error exporting individual report to Excel: {e}")
        return False
def get_activity_log():
    logs = []
    conn = None
    try:
        conn = create_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT timestamp, username, action FROM activity_log ORDER BY log_id DESC")
        for row in cursor.fetchall():
            logs.append(list(row))
        return logs
    except Exception as e:
        print(f"Error reading activity log: {e}")
        return []
    finally:
        if conn:
            conn.close()
def bulk_import_students(filepath):
    logs = {"success": 0, "duplicates": 0, "errors": 0}
    conn = None
    try:
        wb = openpyxl.load_workbook(filepath)
        ws = wb.active
        conn = create_connection()
        cursor = conn.cursor()
        existing_rolls = set()
        cursor.execute("SELECT roll_no FROM students")
        for row in cursor.fetchall():
            existing_rolls.add(str(row[0]))
        for row in ws.iter_rows(min_row=2, values_only=True):
            try:
                if row[0] is None or row[1] is None:
                    logs["errors"] += 1
                    continue 
                roll_no = str(row[0]).strip()
                name = str(row[1]).strip()
                department = str(row[2]).strip() if row[2] else ""
                class_name = str(row[3]).strip() if row[3] else ""
                prn_no = str(row[4]).strip() if row[4] else ""
                contact_no = str(row[5]).strip() if row[5] else ""
                dob = str(row[6]).strip() if row[6] else ""
                if not roll_no or not name:
                    logs["errors"] += 1
                    continue
                if roll_no in existing_rolls:
                    logs["duplicates"] += 1
                    continue 
                cursor.execute("""
                INSERT INTO students (roll_no, name, department, class_name, prn_no, contact_no, dob)
                VALUES (?, ?, ?, ?, ?, ?, ?)
                """, (roll_no, name, department, class_name, prn_no, contact_no, dob))
                existing_rolls.add(roll_no)
                logs["success"] += 1
            except Exception as e:
                print(f"Error processing row {row}: {e}")
                logs["errors"] += 1
        conn.commit()
        return logs
    except Exception as e:
        if conn:
            conn.rollback()
        raise Exception(f"Bulk import failed: {e}")
    finally:
        if conn:
            conn.close()
def get_student_photo_path(roll_no, name):
    try:
        gui_folder = DATASET_DIR / f"{roll_no}_{name}"
        if gui_folder.exists():
            images = list(gui_folder.glob("*.png"))
            if images:
                return str(images[0])
        sanitized_name = name.replace(' ', '_')
        sanitized_name = re.sub(r'[^\w\-]', '', sanitized_name)
        console_folder = STUDENTS_DIR / sanitized_name
        if console_folder.exists():
            images = list(console_folder.glob("*.png"))
            if images:
                return str(images[0])
        return None
    except Exception as e:
        print(f"Error finding photo path: {e}")
        return None
def clear_all_attendance_data():
    conn = None
    try:
        conn = create_connection()
        cursor = conn.cursor()
        cursor.execute("DELETE FROM attendance_logs")
        cursor.execute("DELETE FROM sqlite_sequence WHERE name='attendance_logs'")
        conn.commit()
        return True
    except Exception as e:
        if conn:
            conn.rollback()
        raise Exception(f"Attendance data clear karne mein error: {e}")
    finally:
        if conn:
            conn.close()
def clear_unknown_logs():
    try:
        count = 0
        for a_file in UNKNOWN_LOGS_DIR.glob("*.png"):
            a_file.unlink()
            count += 1
        return count
    except Exception as e:
        raise Exception(f"Unknown logs clear karne mein error: {e}")
def get_monthly_attendance_trend(department=None):
    conn = None
    try:
        conn = create_connection()
        cursor = conn.cursor()
        query_total = "SELECT COUNT(*) FROM students"
        params_total = []
        if department:
            query_total += " WHERE department = ?"
            params_total.append(department)
        cursor.execute(query_total, params_total)
        total_students = cursor.fetchone()[0]
        if total_students == 0:
            return pd.Series(dtype=float) 
        today = datetime.now().date()
        start_date = today - timedelta(days=29)
        query = f"""
        SELECT T1.date, COUNT(DISTINCT T1.roll_no) as present_count
        FROM attendance_logs AS T1
        INNER JOIN students AS T2 ON T1.roll_no = T2.roll_no
        WHERE T1.date BETWEEN ? AND ?
        """
        params = [str(start_date), str(today)]
        if department:
            query += " AND T2.department = ?"
            params.append(department)
        query += " GROUP BY T1.date"
        df = pd.read_sql_query(query, conn, params=params)
        if df.empty:
            return pd.Series(dtype=float)
        df['percentage'] = (df['present_count'] / total_students) * 100
        df['date'] = pd.to_datetime(df['date'])
        df = df.set_index('date')['percentage']
        date_range = pd.date_range(start=start_date, end=today)
        df = df.reindex(date_range, fill_value=0)
        return df
    except Exception as e:
        print(f"Error getting attendance trend: {e}")
        return pd.Series(dtype=float)
    finally:
        if conn:
            conn.close()

# --- YEH NAYA FUNCTION ADD KAREIN ---
def rebuild_all_encodings():
    """
    encodings.pkl ko delete karta hai aur database se sabhi students ki sabhi photos
    (dataset/ aur students/ folders se) ko re-encode karke nayi file banata hai.
    """
    print("Re-building all encodings...")
    
    # 1. Saare students ki list DB se nikalo
    all_students = get_student_details_dict() # Yeh {roll_no: {details}} return karta hai
    if not all_students:
        raise Exception("Database mein koi student nahi hai.")

    new_encodings = []
    new_data = []
    processed_students = 0
    processed_images = 0

    # 2. Har student ke liye loop chalao
    for roll_no, details in all_students.items():
        name = details.get("name")
        if not name:
            continue

        student_image_paths = []
        
        # Path 1 (GUI folder)
        gui_folder = DATASET_DIR / f"{roll_no}_{name}"
        if gui_folder.exists():
            student_image_paths.extend(list(gui_folder.glob("*.png")))

        # Path 2 (Console folder)
        sanitized_name = name.replace(' ', '_')
        sanitized_name = re.sub(r'[^\w\-]', '', sanitized_name)
        console_folder = STUDENTS_DIR / sanitized_name
        if console_folder.exists():
            student_image_paths.extend(list(console_folder.glob("*.png")))
            
        if not student_image_paths:
            print(f"Warning: Student {roll_no} {name} ke liye koi image nahi mili, skip kar rahe hain.")
            continue

        student_encoded = False
        # 3. Har student ki har image ko process karo
        for img_path in student_image_paths:
            try:
                img_rgb = face_recognition.load_image_file(str(img_path))
                boxes = face_recognition.face_locations(img_rgb, model='hog')
                encodings = face_recognition.face_encodings(img_rgb, boxes)

                if encodings:
                    new_encodings.append(encodings[0])
                    new_data.append({"roll": roll_no, "name": name})
                    processed_images += 1
                    student_encoded = True
            except Exception as e:
                print(f"Error processing image {img_path}: {e}")
        
        if student_encoded:
            processed_students += 1
            
    # 4. Nayi encodings file save karo
    if not new_encodings:
        raise Exception("Koi bhi face encode nahi ho saka.")
        
    save_encodings(new_encodings, new_data)
    
    return processed_students, processed_images